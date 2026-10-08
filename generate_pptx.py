import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

DARK_BG = RGBColor(15, 23, 42)      # Slate 900
EMERALD = RGBColor(16, 185, 129)    # Emerald 500
ACCENT_GREEN = RGBColor(5, 150, 105)# Emerald 600
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(203, 213, 225)# Slate 300
CARD_BG = RGBColor(30, 41, 59)      # Slate 800

def set_slide_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category_text="PHARMACY MANAGEMENT SYSTEM"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = EMERALD
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE

# Slide 1: Title Slide
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1, DARK_BG)

title_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.0))
tf = title_box.text_frame
tf.word_wrap = True

p_pre = tf.paragraphs[0]
p_pre.text = "FINAL CAPSTONE ORAL DEFENSE • DELIVERABLE 4"
p_pre.font.size = Pt(14)
p_pre.font.bold = True
p_pre.font.color.rgb = EMERALD
p_pre.space_after = Pt(14)

p_main = tf.add_paragraph()
p_main.text = "Pharmacy Inventory & Sales\nManagement System"
p_main.font.size = Pt(44)
p_main.font.bold = True
p_main.font.color.rgb = WHITE
p_main.space_after = Pt(18)

p_sub = tf.add_paragraph()
p_sub.text = "A Layered, Adversarial-Resistant Full-Stack Web Architecture | Python • Flask • SQLite • Vanilla JS"
p_sub.font.size = Pt(16)
p_sub.font.color.rgb = LIGHT_GRAY

footer_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.8))
p_f = footer_box.text_frame.paragraphs[0]
p_f.text = "Presented by Team 5 | Final Term Capstone Defense"
p_f.font.size = Pt(13)
p_f.font.color.rgb = LIGHT_GRAY

# Slide 2: The Problem
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2, DARK_BG)
add_header(s2, "The Real-World Retail Pharmacy Problem")

problems = [
    ("Uncontrolled Stock Leaks & Spoilage", "Manual paper records and isolated spreadsheets fail to flag expiring medicine batches, leading to unexpected financial loss and regulatory liability."),
    ("Slow Checkout & Reconciliation", "Unsynchronized customer records, manual tax calculations, and disjointed invoices create bottlenecks and reconciliation errors during peak hours."),
    ("Fragile Software Vulnerabilities", "Off-the-shelf basic systems crash on edge inputs (negative balances, arithmetic overflows) and lack defense-in-depth against common script injections (XSS).")
]

for i, (title, desc) in enumerate(problems):
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.9), Inches(2.0), Inches(3.7), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = EMERALD
    card.line.width = Pt(1)
    
    tf = card.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = f"0{i+1}"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD
    p1.space_after = Pt(12)
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(13)
    p3.font.color.rgb = LIGHT_GRAY

# Slide 3: The Solution
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3, DARK_BG)
add_header(s3, "The Solution: A Unified, Layered Monolith")

solutions = [
    ("Full Lifecycle Pharmacy Suite", "Covers end-to-end workflows: Catalog, Customers, Suppliers, Invoicing, Stock Adjustments, Purchases, and Live Financial Reports."),
    ("Asynchronous, Zero-Reload UI", "Vanilla JavaScript fetch integrations ensure smooth workflows with immediate inline field validations and instant loading feedback."),
    ("Security & Guarded Core", "Mandatory session protection (@app.before_request), route validation guards, and SQL-overflow / XSS sanitization.")
]

for i, (title, desc) in enumerate(solutions):
    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.9), Inches(2.0), Inches(3.7), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = ACCENT_GREEN
    card.line.width = Pt(1)
    
    tf = card.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = f"FEATURE 0{i+1}"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD
    p1.space_after = Pt(10)
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(13)
    p3.font.color.rgb = LIGHT_GRAY

# Slide 4: System Architecture
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4, DARK_BG)
add_header(s4, "System Architecture & Request-Response Pipeline")

arch_steps = [
    ("1. Client View Layer", "Semantic HTML5 + CSS + Jinja2 templates. JS intercepts submits, displays spinners, and prevents page reload."),
    ("2. Routing & Guard Tier", "Flask Blueprints. Enforces session auth and inspects incoming payloads with strict validation guard clauses."),
    ("3. Controller & SQL Tier", "Decoupled Python functions execute explicit, parameterized SQLite queries with foreign keys and ACID integrity.")
]

for i, (title, desc) in enumerate(arch_steps):
    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.9), Inches(2.0), Inches(3.7), Inches(3.0))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = EMERALD
    
    tf = card.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(17)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(10)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(13)
    p2.font.color.rgb = LIGHT_GRAY

env_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.3), Inches(11.5), Inches(1.5))
env_box.fill.solid()
env_box.fill.fore_color.rgb = CARD_BG
env_box.line.color.rgb = ACCENT_GREEN
tf_e = env_box.text_frame
tf_e.word_wrap = True
pe1 = tf_e.paragraphs[0]
pe1.text = "Standardized JSON Response Contract"
pe1.font.size = Pt(15)
pe1.font.bold = True
pe1.font.color.rgb = EMERALD
pe2 = tf_e.add_paragraph()
pe2.text = "• Success: {'status': 200/201, 'data': {...}}   |   • Validation Fail: {'status': 422, 'error': 'message', 'field': 'key'}\n• Resource Missing: {'status': 404, 'error': 'Not found'}   |   • Server Guard: Never exposes raw 500 stack traces."
pe2.font.size = Pt(13)
pe2.font.color.rgb = LIGHT_GRAY

# Slide 5: Live Demo Click-Path
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5, DARK_BG)
add_header(s5, "Rehearsed Live Demo Flow (Rubric-Aligned)")

demo_steps = [
    ("1. Secure Authentication", "Demonstrate session guard redirect to /login. Log in as admin / admin123."),
    ("2. Real-Time Dashboard", "Show dynamic aggregation widgets (Medicines, Stock Value, Sales, Customers)."),
    ("3. Graceful Failure (P0 Demo)", "Intentionally submit a medicine with negative price or blank name. Show inline 422 error without page reload."),
    ("4. Happy Path Creation", "Add 'Amoxicillin 500mg' successfully. View instant appearance in catalog."),
    ("5. Dynamic Inventory Adjustment", "Adjust quantity on a lot. Demonstrate real-time inventory recalculation."),
    ("6. Financial Reports", "Review automated revenue, purchasing expenditure, and net profit margins.")
]

for i, (title, desc) in enumerate(demo_steps):
    col = i % 2
    row = i // 2
    box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col*5.9), Inches(1.8 + row*1.7), Inches(5.6), Inches(1.4))
    box.fill.solid()
    box.fill.fore_color.rgb = CARD_BG
    box.line.color.rgb = EMERALD
    
    tf = box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(12)
    p2.font.color.rgb = LIGHT_GRAY

# Slide 6: Team Roles & Code Defense
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6, DARK_BG)
add_header(s6, "5-Member Role Assignment & Defense Ownership")

roles = [
    ("Member 1: Architecture & Auth", "app.py, routes/extras.py", "System architecture, session guard (@app.before_request), and login workflow."),
    ("Member 2: Medicines & Validation", "routes/medicines.py, controllers/medicines.py", "Medicines CRUD, XSS prevention guard, and 32-bit integer overflow protection."),
    ("Member 3: Sales & Transactions", "routes/sales.py, routes/customers.py", "Invoicing, customer linkages, financial amount validation, and data persistence."),
    ("Member 4: Procurement & Inventory", "controllers/suppliers.py, routes/extras.py", "Supplier directory, active stock adjustments (UPDATE current_stock), and reports."),
    ("Member 5: Frontend & Automation QA", "static/css/style.css, templates/*, tests/*", "Async JS fetch lifecycles, inline 422 error painting, and pytest automated test suite.")
]

for i, (name, files, desc) in enumerate(roles):
    box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + i*1.0), Inches(11.5), Inches(0.85))
    box.fill.solid()
    box.fill.fore_color.rgb = CARD_BG
    box.line.color.rgb = EMERALD
    
    tf = box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = f"{name}  |  Files: {files}"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.text = f"Defense Focus: {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = LIGHT_GRAY

# Slide 7: QA & Adversarial Hardening
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7, DARK_BG)
add_header(s7, "Quality Assurance & Adversarial Hardening")

qa_items = [
    ("Adversarial QA Discovery (Week 10)", "During feature freeze testing, the system was subjected to hostile inputs. Two critical issues surfaced: script injection (<script>) and database integer overflow when stock numbers exceeded 32-bit maximums."),
    ("Hardened Route Guards (Week 11)", "Validation guards were refactored to explicitly block <script> tokens and cap stock at 2,147,483,647. Both return safe 422 HTTP responses rather than allowing server crashes."),
    ("100% Automated Regression Suite", "A comprehensive 10-test automated pytest suite runs against isolated database runners, verifying both happy-path transactions and adversarial attacks.")
]

for i, (title, desc) in enumerate(qa_items):
    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.9), Inches(2.0), Inches(3.7), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = EMERALD
    card.line.width = Pt(1)
    
    tf = card.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = f"PILLAR 0{i+1}"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD
    p1.space_after = Pt(10)
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(13)
    p3.font.color.rgb = LIGHT_GRAY

# Slide 8: Key Lessons & Conclusion
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8, DARK_BG)
add_header(s8, "What We Learned & Ready for Defense")

lessons = [
    ("Guard Clauses Save Systems", "Placing strict validation at the route boundary before touching controllers prevents unexpected exceptions from reaching the database layer."),
    ("Uniform Response Envelopes", "Adopting standardized {status, data, error} structures across all endpoints made frontend error binding clean and predictable."),
    ("AI as an Accelerator with Vigilance", "AI tools rapidly draft boilerplate code, but thorough human review, adversarial testing, and code audits are required to catch subtle security flaws.")
]

for i, (title, desc) in enumerate(lessons):
    card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*3.9), Inches(2.0), Inches(3.7), Inches(3.8))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = ACCENT_GREEN
    
    tf = card.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = f"TAKEAWAY {i+1}"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD
    p1.space_after = Pt(8)
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(13)
    p3.font.color.rgb = LIGHT_GRAY

fin_box = s8.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.5), Inches(0.8))
p_fin = fin_box.text_frame.paragraphs[0]
p_fin.text = "Thank you! We are ready for the unassisted oral defense and technical questions."
p_fin.font.size = Pt(18)
p_fin.font.bold = True
p_fin.font.color.rgb = EMERALD
p_fin.alignment = PP_ALIGN.CENTER

prs.save("Presentation.pptx")
print("Presentation.pptx generated successfully!")
